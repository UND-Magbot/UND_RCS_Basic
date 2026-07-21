"""
LiDAR 감지 → 랙 인식·이송 Flow 문서(.docx) 생성기

다이어그램은 matplotlib 으로 PNG 렌더 후 Word 에 삽입한다.
(mermaid-cli 는 Chromium 다운로드가 필요해 오프라인 환경에서 부적합)

실행:
    python docs/generate_rack_flow_docx.py
산출물:
    docs/RCS-Basic_랙인식_이송_Flow.docx
    docs/img/flow_*.png
"""
from __future__ import annotations

import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon, FancyArrowPatch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

# ── 한글 폰트 ──
plt.rcParams["font.family"] = ["Malgun Gothic", "AppleGothic", "NanumGothic", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

BASE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE, "img")
os.makedirs(IMG_DIR, exist_ok=True)

# ── 색상 팔레트 ──
C_START = "#E8EAED"   # 시작/종료
C_PROC = "#DCE9F7"    # 처리
C_API = "#D6EFDC"     # 로봇 API 호출
C_DEC = "#FCEFC7"     # 판정
C_FAIL = "#F8D7DA"    # 실패
EDGE = "#5F6368"
TXT = "#202124"


# ── 도형 헬퍼 ──

def box(ax, x, y, w, h, text, color=C_PROC, fs=8.5, rounded=0.08):
    """중심 (x, y) 기준 사각형 노드"""
    ax.add_patch(FancyBboxPatch(
        (x - w / 2, y - h / 2), w, h,
        boxstyle=f"round,pad=0.02,rounding_size={rounded}",
        linewidth=1.1, edgecolor=EDGE, facecolor=color, zorder=2,
    ))
    ax.text(x, y, text, ha="center", va="center", fontsize=fs, color=TXT, zorder=3, linespacing=1.4)
    return (x, y, w, h)


def diamond(ax, x, y, w, h, text, color=C_DEC, fs=8):
    """판정 노드 (마름모)"""
    pts = [(x, y + h / 2), (x + w / 2, y), (x, y - h / 2), (x - w / 2, y)]
    ax.add_patch(Polygon(pts, closed=True, linewidth=1.1,
                         edgecolor=EDGE, facecolor=color, zorder=2))
    ax.text(x, y, text, ha="center", va="center", fontsize=fs, color=TXT, zorder=3, linespacing=1.3)
    return (x, y, w, h)


def arrow(ax, p1, p2, label="", style="-|>", color=EDGE, rad=0.0, fs=7.5, lx=0, ly=0):
    """p1 → p2 화살표. p1/p2 는 (x, y) 좌표."""
    ax.add_patch(FancyArrowPatch(
        p1, p2, arrowstyle=style, mutation_scale=11,
        linewidth=1.0, color=color, zorder=1,
        connectionstyle=f"arc3,rad={rad}",
    ))
    if label:
        mx, my = (p1[0] + p2[0]) / 2 + lx, (p1[1] + p2[1]) / 2 + ly
        ax.text(mx, my, label, ha="center", va="center", fontsize=fs,
                color="#B06000", zorder=4,
                bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.9))


def new_ax(w_in, h_in, xlim, ylim):
    fig, ax = plt.subplots(figsize=(w_in, h_in), dpi=200)
    ax.set_xlim(*xlim)
    ax.set_ylim(*ylim)
    ax.axis("off")
    return fig, ax


def save(fig, name):
    path = os.path.join(IMG_DIR, name)
    fig.savefig(path, bbox_inches="tight", pad_inches=0.15, facecolor="white")
    plt.close(fig)
    print(f"  saved: {path}")
    return path


# ── 다이어그램 1: 전체 흐름 ──

def draw_overall():
    fig, ax = new_ax(9.6, 15.0, (0, 11.6), (0, 34))

    box(ax, 5, 33, 3.4, 1.2, "작업 시작 (run_route_job)", C_START)
    box(ax, 5, 31, 5.0, 1.5, "위치 보정\nPOST /services/start_global_positioning\nuse_barcode + use_base_map_match", C_API)
    diamond(ax, 5, 28.7, 3.6, 1.9, "WS /global_positioning_state\nsucceeded ?")
    box(ax, 8.6, 28.7, 2.4, 1.3, "current-map 재선택\n후 1회 재시도", C_PROC, fs=8)

    box(ax, 5, 26.2, 5.0, 1.4, "랙 위치 POI(standby)로 랙 정렬\nPOST /chassis/moves  type=align_with_rack", C_API)
    diamond(ax, 5, 24.0, 3.0, 1.6, "state ?")
    box(ax, 8.5, 24.0, 2.6, 1.3, "재시도 루틴\n(다이어그램 3)", C_FAIL, fs=8)
    box(ax, 8.9, 21.6, 2.6, 1.2, "작업 실패\ncancel + jack_down", C_FAIL, fs=8)

    box(ax, 5, 21.8, 4.4, 1.2, "POST /services/jack_up  + 10초 대기", C_API)
    box(ax, 5, 19.8, 4.8, 1.3, "작업 위치로 랙 이송\ntype=to_unload_point", C_API)
    diamond(ax, 5, 17.6, 3.0, 1.6, "state ?")
    box(ax, 1.5, 17.6, 2.6, 1.3, "safe_move 재시도\n5초 간격 · 최대 200회", C_PROC, fs=8)

    box(ax, 5, 15.4, 4.4, 1.2, "POST /services/jack_down  + 10초 대기", C_API)
    diamond(ax, 5, 13.2, 3.2, 1.7, "수동 확인 모드 ?")
    box(ax, 8.4, 13.2, 2.8, 1.2, "출발 버튼 대기\n(최대 30분)", C_PROC, fs=8)
    box(ax, 1.6, 13.2, 2.8, 1.2, "wait_sec 대기", C_PROC, fs=8)

    box(ax, 5, 10.9, 4.6, 1.3, "랙 재정렬 (align_with_rack)\n→ jack_up", C_API)
    diamond(ax, 5, 8.6, 3.4, 1.7, "다음 목적지 있음 ?")

    box(ax, 5, 6.2, 4.8, 1.3, "랙 위치(standby)로 복귀\nto_unload_point → jack_down", C_API)
    box(ax, 5, 4.0, 5.0, 1.4, "충전소 복귀\nstandard 접근 → type=charge\ncharge_retry_count=3", C_API)
    box(ax, 5, 1.9, 2.6, 1.1, "완료", C_START)

    # 화살표
    arrow(ax, (5, 32.4), (5, 31.8))
    arrow(ax, (5, 30.25), (5, 29.65))
    arrow(ax, (6.8, 28.7), (7.4, 28.7), "실패")
    arrow(ax, (8.6, 28.05), (5.6, 27.0), rad=-0.15)
    arrow(ax, (5, 27.75), (5, 26.9), "성공", lx=0.6)
    arrow(ax, (5, 25.5), (5, 24.8))
    arrow(ax, (6.5, 24.0), (7.2, 24.0), "failed")
    arrow(ax, (8.5, 23.35), (8.85, 22.2), "최종 실패", lx=1.0)
    arrow(ax, (7.2, 23.6), (6.4, 22.4), "성공", rad=0.2, lx=-0.3)   # 재시도 성공 → jack_up
    arrow(ax, (5, 23.2), (5, 22.4), "succeeded", lx=0.9)
    arrow(ax, (5, 21.2), (5, 20.45))
    arrow(ax, (5, 19.15), (5, 18.4))
    arrow(ax, (3.5, 17.6), (2.8, 17.6), "failed / timeout", ly=0.35)
    arrow(ax, (1.5, 18.25), (4.0, 19.4), rad=0.15)
    arrow(ax, (5, 16.8), (5, 16.0), "succeeded", lx=0.85)
    arrow(ax, (5, 14.8), (5, 14.05))
    arrow(ax, (6.6, 13.2), (7.0, 13.2), "예")
    arrow(ax, (3.4, 13.2), (3.0, 13.2), "아니오")
    arrow(ax, (8.4, 12.6), (5.8, 11.55), rad=-0.12)
    arrow(ax, (1.6, 12.6), (4.2, 11.55), rad=0.12)
    arrow(ax, (5, 10.25), (5, 9.45))
    # 루프백: 다음 목적지 있음 → 이송 단계로 (우측 세로 라인)
    arrow(ax, (6.7, 8.6), (10.8, 8.6), "예", ly=0.35, style="-")
    arrow(ax, (10.8, 8.6), (10.8, 19.8), style="-")
    arrow(ax, (10.8, 19.8), (7.45, 19.8))
    arrow(ax, (5, 7.75), (5, 6.85), "아니오", lx=0.8)
    arrow(ax, (5, 5.55), (5, 4.7))
    arrow(ax, (5, 3.3), (5, 2.45))

    ax.text(0.1, 33.6, "[1] 전체 흐름 — rack_pickup 모드", fontsize=12, weight="bold", color=TXT)
    return save(fig, "flow_overall.png")


# ── 다이어그램 2: align_with_rack 내부 (LiDAR 랙 인식) ──

def draw_align():
    fig, ax = new_ax(8.6, 12.0, (0, 10), (0, 26))

    box(ax, 4.6, 25.0, 5.6, 1.4, "POST /chassis/moves\ntype=align_with_rack, target = Shelves Point(type=34)", C_START, fs=8.5)
    box(ax, 4.6, 22.8, 4.6, 1.1, "Shelves Point 근처까지 일반 경로 주행", C_PROC)
    box(ax, 4.6, 20.7, 5.0, 1.3, "LiDAR 스캔 — 점군에서\n랙 다리(leg) 후보 추출", C_API)
    diamond(ax, 4.6, 18.1, 5.4, 2.4, "rack.specs 와 매칭 ?\nleg_size · foot_radius · width × depth")
    box(ax, 4.6, 15.3, 5.0, 1.3, "랙 4다리 중심점 계산\n→ 랙 중심 · 방향(yaw) 산출", C_PROC)
    box(ax, 4.6, 13.2, 4.4, 1.1, "중심선 정렬 (alignment=center)", C_PROC)
    box(ax, 4.6, 11.1, 5.0, 1.3, "랙 하부로 저속 진입\nmargin 0.05~0.1 m 유지", C_PROC)
    diamond(ax, 4.6, 8.3, 5.6, 2.4, "잭 전면 여유 확인\ncargo_to_jack_front_edge_min_distance ≥ 0.05 ?")
    box(ax, 4.6, 5.3, 4.8, 1.4, "state = succeeded\nWS /detected_rack 갱신", C_START)

    box(ax, 8.6, 13.2, 2.6, 1.6, "fail_message\nfailed to find rack /\nAlignFailedInRackArea", C_FAIL, fs=7.5)
    box(ax, 8.6, 10.0, 2.2, 1.0, "state = failed", C_FAIL)

    arrow(ax, (4.6, 24.3), (4.6, 23.35))
    arrow(ax, (4.6, 22.25), (4.6, 21.35))
    arrow(ax, (4.6, 20.05), (4.6, 19.3))
    arrow(ax, (7.3, 18.1), (8.6, 18.1), "매칭 실패")
    arrow(ax, (8.6, 18.1), (8.6, 14.0), style="-|>")
    arrow(ax, (4.6, 16.9), (4.6, 15.95), "매칭 성공", lx=0.85)
    arrow(ax, (4.6, 14.65), (4.6, 13.75))
    arrow(ax, (4.6, 12.65), (4.6, 11.75))
    arrow(ax, (4.6, 10.45), (4.6, 9.5))
    arrow(ax, (7.4, 8.3), (7.6, 9.6), "여유 부족", rad=-0.15, lx=0.75, ly=-0.35)
    arrow(ax, (8.6, 12.4), (8.6, 10.6))
    arrow(ax, (4.6, 7.1), (4.6, 6.0), "충분", lx=0.6)

    ax.text(0.0, 25.9, "[2] align_with_rack 내부 — LiDAR 랙 인식 구간", fontsize=12, weight="bold", color=TXT)
    ax.text(0.0, 3.6, "모니터링 WebSocket 토픽", fontsize=9.5, weight="bold", color=TXT)
    ax.text(0.0, 2.6,
            "/detected_rack — 랙 감지 여부 및 좌표\n"
            "/jack_state — jacking_up / jacking_down / hold\n"
            "/robot_model — 잭 업 시 풋프린트가 랙 크기로 확장",
            fontsize=8.5, color=TXT, linespacing=1.6, va="top")
    ax.text(0.0, 0.2, "※ 다리 후보 추출 → spec 매칭 → 중심 산출 세부 단계는 펌웨어 내부 동작으로,\n"
                      "   rack.specs 필드 의미와 실패 메시지로부터 역추정한 내용임.",
            fontsize=7.5, color="#A05000", linespacing=1.5, va="bottom")
    return save(fig, "flow_align.png")


# ── 다이어그램 3: 랙 인식 실패 재시도 ──

def draw_retry():
    fig, ax = new_ax(11.5, 4.6, (0, 25.2), (0, 8))

    steps = [
        (2.5, "1차\n그대로 시도"),
        (7.9, "2차\n재로컬화\n(start_global_positioning)"),
        (13.3, "3차\n0.5 m 후진 → 재로컬화\n(LiDAR 시야 확보)"),
        (18.7, "4차\n측면 0.3 m 우회\n→ 재로컬화\n(접근 각도 변경)"),
    ]
    for x, label in steps:
        box(ax, x, 5.6, 4.4, 2.1, label, C_PROC, fs=8.0)

    box(ax, 23.3, 5.6, 3.0, 1.3, "최종 실패 반환", C_FAIL, fs=8)
    box(ax, 10.6, 1.4, 4.6, 1.2, "succeeded — 잭 업 진행", C_START)

    for i in range(3):
        arrow(ax, (steps[i][0] + 2.2, 5.6), (steps[i + 1][0] - 2.2, 5.6), "실패", ly=0.5)
    arrow(ax, (20.9, 5.6), (21.8, 5.6), "실패", ly=0.5)

    # 각 시도에서 성공하면 잭 업 단계로 수렴
    for i, (x, _) in enumerate(steps):
        arrow(ax, (x, 4.55), (10.6 + (i - 1.5) * 0.9, 2.05),
              "성공" if i == 0 else "", color="#1E8E3E", fs=8, lx=-0.7)

    ax.text(0.0, 7.6, "[3] 랙 인식 실패 시 재시도 — align_with_retry (최대 4회)",
            fontsize=12, weight="bold", color=TXT)
    return save(fig, "flow_retry.png")


# ── Word 문서 ──

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        run.font.name = "맑은 고딕"
    return h


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    return t


def add_image(doc, path, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))


def build_docx(img_overall, img_align, img_retry):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    # 표지
    title = doc.add_heading("LiDAR 감지 → 랙(선반) 인식 · 이송 Flow", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("RCS-Basic × AutoXing crawler (S300 / S600)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref = doc.add_paragraph("근거: AutoXing AxBot REST Book + BackEnd/app/services/jack_service.py")
    ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in ref.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    # 1. 전제 조건
    add_heading(doc, "1. 전제 조건", 1)
    doc.add_paragraph("아래 세 가지 중 하나라도 어긋나면 로봇은 랙을 인식하지 못한다 "
                      "(failed to find rack). 랙 인식 실패 시 가장 먼저 점검할 항목이다.")
    add_table(doc,
              ["항목", "설정 위치", "값"],
              [
                  ["랙 물리 스펙", "PATCH /system/settings/user → rack.specs",
                   "S300: 0.765 × 0.765 / S600: 0.83 × 0.87\n(app/constants/rack_specs.py)"],
                  ["Shelves Point", "맵 overlay feature",
                   'type="34", subtype="rack",\nhasFixedLegs=false, dockViaDirection="front"'],
                  ["overlay 반영", "overlay PATCH 직후",
                   "POST /chassis/current-map 재선택\n(overlays_version 갱신)"],
              ],
              widths=[3.0, 6.0, 7.0])
    p = doc.add_paragraph()
    r = p.add_run("※ rack.specs 는 실제 사용하는 사이즈 종류만 등록할 것. spec 이 여러 개면 펌웨어가 "
                  "LiDAR 측정값을 엉뚱한 spec 에 매칭해 실패하는 사례가 관찰되었다.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xA0, 0x50, 0x00)

    # 2. 전체 흐름
    doc.add_page_break()
    add_heading(doc, "2. 전체 흐름 (rack_pickup 모드)", 1)
    doc.add_paragraph("위치 보정 → 랙 정렬 → 잭 업 → 이송 → 잭 다운 → 랙 위치 복귀 → 충전소 도킹.")
    add_image(doc, img_overall, width_cm=15.0)

    # 3. align 내부
    doc.add_page_break()
    add_heading(doc, "3. align_with_rack 내부 — LiDAR 랙 인식 구간", 1)
    doc.add_paragraph("LiDAR 점군에서 랙 다리를 검출해 rack.specs 와 매칭하고, 4다리의 중심을 "
                      "계산해 랙 중심선에 정렬한 뒤 랙 하부로 진입한다.")
    add_image(doc, img_align, width_cm=14.5)

    # 4. 재시도
    doc.add_page_break()
    add_heading(doc, "4. 랙 인식 실패 시 재시도 (align_with_retry)", 1)
    doc.add_paragraph("한 번의 인식 실패로 작업을 중단하지 않고, 실패 원인별로 회복 동작을 "
                      "단계적으로 바꿔 최대 4회까지 재시도한다.")
    add_image(doc, img_retry, width_cm=16.5)

    add_heading(doc, "4.1 실패 코드별 원인과 대응", 2)
    add_table(doc,
              ["실패 코드", "원인", "대응"],
              [
                  ["failed to find rack", "POI type 이 34 가 아님 / 랙 부재 / spec 불일치",
                   "overlay · rack.specs 점검"],
                  ["AlignFailedInRackArea", "진입 각도 · LiDAR 시야 문제", "후진 · 측면 우회 재시도"],
                  ["NoFreeSpaceInRackArea", "랙 영역에 빈 공간 없음", "대상 영역 확인"],
                  ["MoveTimeout", "경로 막힘", "safe_move 자동 재시도 (5초 간격, 최대 200회)"],
                  ["ChargeRetryCountExceeded", "충전 도킹 실패",
                   "사전 접근 POI (<충전소명>-1) 경유 후 도킹"],
              ],
              widths=[4.0, 6.0, 6.0])

    # 5. work_mode
    add_heading(doc, "5. work_mode 별 분기", 1)
    add_table(doc,
              ["모드", "랙 인식 (align_with_rack)", "이송 move type", "잭 조작"],
              [
                  ["rack_pickup", "사용 — LiDAR 랙 정렬", "to_unload_point",
                   "pickup 에서 up, dropoff 에서 down"],
                  ["delivery_no_rack", "미사용", "standard", "pickup up / dropoff down"],
                  ["simple_move", "미사용", "standard", "없음"],
              ],
              widths=[3.2, 4.8, 3.5, 4.5])

    out = os.path.join(BASE, "RCS-Basic_랙인식_이송_Flow.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    print("다이어그램 렌더링...")
    i1 = draw_overall()
    i2 = draw_align()
    i3 = draw_retry()
    print("Word 문서 생성...")
    out = build_docx(i1, i2, i3)
    print(f"완료: {out}")

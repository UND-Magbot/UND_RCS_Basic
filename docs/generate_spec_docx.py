# -*- coding: utf-8 -*-
"""
RCS-Basic 시스템 사양서 (.docx) 생성기.
실행: python docs/generate_spec_docx.py
출력: docs/RCS-Basic_시스템_사양서.docx
"""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUTPUT = Path(__file__).resolve().parent / "RCS-Basic_시스템_사양서.docx"

KOR_FONT = "맑은 고딕"
MONO_FONT = "Consolas"

PRIMARY = RGBColor(0x1F, 0x4E, 0x79)   # 진한 파랑
SUB = RGBColor(0x2E, 0x75, 0xB6)
MUTED = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0xC0, 0x50, 0x4D)
SUCCESS = RGBColor(0x2E, 0x7D, 0x32)


# ──────────────────────────────────────────────────────────
# 폰트 / 스타일 유틸
# ──────────────────────────────────────────────────────────
def _set_font(run, name=KOR_FONT, size=10.5, bold=False, color=None, mono=False):
    if mono:
        name = MONO_FONT
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # 한글 폰트 적용
    rPr = run._element.rPr
    if rPr is None:
        rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    _set_font(r, size=20, bold=True, color=PRIMARY)


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    _set_font(r, size=15, bold=True, color=SUB)


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    _set_font(r, size=12.5, bold=True, color=MUTED)


def para(doc, text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.runs[0] if p.runs else p.add_run("")
    r.text = text
    _set_font(r, size=10.5)


def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    _set_font(r, size=9.5, mono=True, color=RGBColor(0x33, 0x33, 0x33))
    # 회색 음영
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)


def table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light List Accent 1"
    t.autofit = False

    if col_widths is not None:
        for col, cm in zip(t.columns, col_widths):
            for cell in col.cells:
                cell.width = Cm(cm)

    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        _set_font(r, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            _set_font(r, size=10)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ──────────────────────────────────────────────────────────
# 문서 작성
# ──────────────────────────────────────────────────────────
def build():
    doc = Document()

    # 페이지 여백
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # 표지
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("RCS-Basic 관제 시스템")
    _set_font(r, size=28, bold=True, color=PRIMARY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("기술 사양서")
    _set_font(r, size=18, bold=True, color=SUB)

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = info.add_run(
        f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}\n"
        "프로젝트: RCS-Basic (Robot Control System – Basic)\n"
        "대상 로봇: AutoXing crawler_s300_op5 (펌웨어 2.12.21-opi64)"
    )
    _set_font(r, size=11, color=MUTED)

    doc.add_page_break()

    # ───────────────────────────────────────────────
    h1(doc, "1. 시스템 개요")
    para(
        doc,
        "RCS-Basic 은 AutoXing 사 모바일 로봇을 로컬망에서 직접 제어하는 경량 관제 시스템이다. "
        "기존 풀버전 RCS 에서 운영 현장에 불필요한 기능(Convoy 대열주행, ACS 화재경보 연동, WCS 인터페이스, "
        "다층 권한 관리, 시스템 로그 분리, 엑셀 내보내기, 작업 이동 관리 등)을 제거하고 "
        "현장 운영에 핵심적인 기능만 남겨 단순화한 형태이다.",
    )
    h3(doc, "운영 환경 가정")
    bullet(doc, "단일 사이트(영역) — 다중 영역/다층 환경도 area_id 단위 운영 가능")
    bullet(doc, "로컬망(LAN) 직결 — 클라우드/CSP 미사용, 외부망 노출 없음")
    bullet(doc, "MySQL/관제 서버 1대, 로봇 N대, 운영자용 PC + 태블릿")
    bullet(doc, "로봇 ↔ 관제 간 통신은 모두 평문 REST/WS (로컬망 신뢰 환경 가정)")

    # ───────────────────────────────────────────────
    h1(doc, "2. 하드웨어 / 로봇 사양")
    h2(doc, "2.1 대상 로봇")
    table(
        doc,
        ["항목", "사양"],
        [
            ["모델", "AutoXing crawler_s300_op5 (기준)"],
            ["펌웨어", "2.12.21-opi64"],
            ["구동 방식", "차동 구동 (Differential drive)"],
            ["최대 전진속도", "사용자 지정 (기본 1.2 m/s, DB robots.max_speed)"],
            ["충전 방식", "자율 도킹 (charge 명령)"],
            ["적재 방식", "리프트 잭(서비스 jack_up/jack_down) 기반"],
            ["적재 형태", "랙(카트) 또는 잭 상판 직접 적재"],
        ],
        col_widths=[5, 11],
    )

    h2(doc, "2.2 센서 (라이다 기반 환경/랙 인식)")
    para(
        doc,
        "AutoXing 펌웨어는 본체 전·후방 2D LiDAR 를 사용해 SLAM, 글로벌 위치추정, "
        "장애물 회피, 그리고 본 시스템의 핵심 기능인 \"랙 다리 인식 기반 정렬\"(align_with_rack) 을 수행한다.",
    )
    bullet(doc, "2D LiDAR (펌웨어 내장) — 매핑(SLAM) + 글로벌 위치추정 + 동적 장애물 회피")
    bullet(doc, "랙 다리 검출 — 정렬 단계에서 등록된 랙 사이즈(rack.specs)와 LiDAR 측정 다리 4점 패턴을 매칭")
    bullet(doc, "근접 로봇 감지 — /nearby_robots 토픽으로 페어링된 다른 로봇 인식 (다중 로봇 양보 트리거)")
    bullet(doc, "잭 상태 감지 — /jack_state (jacking_up / jacking_down / hold)")
    bullet(doc, "랙 감지 상태 — /detected_rack (인식 성공/실패, 다리 좌표)")

    h2(doc, "2.3 잭킹 메커니즘")
    para(
        doc,
        "잭은 본체 상부에 위치한 리프트 플랫폼이다. POI 위치에서 jack_up 시 약 10cm 상승해 "
        "랙(카트)을 들어 올리거나, 상판 위 적재물을 들어 운반한다. jack_down 시 랙/적재물을 내려놓는다. "
        "잭 업/다운에는 약 10초가 소요되며, 일부 펌웨어는 잭 다운 직후 짧은 시간 동안 후속 명령에 400 오류를 반환한다 — "
        "본 시스템은 5초 간격 5회 재시도 + 고정 대기로 이를 보정한다.",
    )

    # ───────────────────────────────────────────────
    h1(doc, "3. 기술 스택")

    h2(doc, "3.1 백엔드")
    table(
        doc,
        ["구분", "기술"],
        [
            ["언어", "Python 3.10"],
            ["웹 프레임워크", "FastAPI 0.121.3"],
            ["ASGI 서버", "uvicorn"],
            ["ORM", "SQLAlchemy 2.0.45 + PyMySQL"],
            ["DB", "MySQL 8.0"],
            ["스케줄러", "APScheduler (BackgroundScheduler)"],
            ["WebSocket 클라이언트", "websocket-client 1.9.0"],
            ["HTTP 클라이언트", "requests"],
            ["동시성", "threading + 자체 락 매니저(POI/Zone)"],
            ["로깅", "stdlib logging + activity_logs 테이블"],
        ],
        col_widths=[5, 11],
    )

    h2(doc, "3.2 프론트엔드 (관제 웹)")
    table(
        doc,
        ["구분", "기술"],
        [
            ["언어", "TypeScript"],
            ["프레임워크", "Next.js 16 (App Router)"],
            ["빌드러", "Turbopack"],
            ["3D 렌더링", "Three.js (모니터링 페이지)"],
            ["스타일", "Vanilla CSS (전역 변수 + 일부 인라인)"],
            ["통신", "fetch (1~2초 폴링)"],
            ["상태", "React useState / useEffect / Custom Event"],
        ],
        col_widths=[5, 11],
    )

    h2(doc, "3.3 태블릿 앱 (운영자 단말)")
    table(
        doc,
        ["구분", "기술"],
        [
            ["플랫폼", "Android (Kotlin)"],
            ["AGP / Gradle", "8.3.2 / 8.7"],
            ["minSdk / targetSdk / compileSdk", "24 / 35 / 35"],
            ["렌더링 방식", "WebView 로 서버 HTML 동적 로드"],
            ["설정 저장", "SharedPreferences (서버 주소, 로봇 ID)"],
            ["두 종류", "TabletApp (수동배차+제어) / TabletAppSimple (단순 표시)"],
        ],
        col_widths=[6, 10],
    )

    h2(doc, "3.4 인프라")
    table(
        doc,
        ["구분", "기술"],
        [
            ["컨테이너", "Docker Compose v2 (권장)"],
            ["호스트 OS", "Ubuntu 22.04"],
            ["DB 호스팅", "호스트 OS 직접 설치 (도커 외부)"],
            ["네트워크", "도커 backend 에 extra_hosts: host.docker.internal:host-gateway"],
            ["포트", "8002 (API) / 3002 (Web) / 3306 (MySQL)"],
            ["헬스체크", "GET /ping, GET /health (DB 연결 확인)"],
        ],
        col_widths=[6, 10],
    )

    # ───────────────────────────────────────────────
    h1(doc, "4. 아키텍처 / 통신 방식")
    h2(doc, "4.1 컴포넌트 구성도 (개념)")
    code_block(
        doc,
        "┌──────────────┐  REST 8002   ┌──────────────┐  REST 8090 / WS  ┌──────────┐\n"
        "│  관제 웹 PC   │ ───────────▶│ Backend (Py) │ ───────────────▶ │ 로봇 N대 │\n"
        "│ (Next.js)    │ ◀───────────│ (FastAPI)    │ ◀─────────────── │ AutoXing│\n"
        "└──────────────┘             └──────┬───────┘                  └──────────┘\n"
        "                                    │ SQL\n"
        "                                    ▼\n"
        "                              ┌──────────────┐\n"
        "                              │ MySQL 8.0    │\n"
        "                              └──────────────┘\n"
        "        ▲ REST 8002 (HTML 폴링)\n"
        "        │\n"
        "┌──────────────┐\n"
        "│  태블릿 앱   │ (WebView → /api/tasks/tablet/{id})\n"
        "└──────────────┘",
    )

    h2(doc, "4.2 통신 채널 요약")
    table(
        doc,
        ["주체 → 대상", "프로토콜", "용도"],
        [
            ["관제 웹 → 백엔드", "REST :8002", "모든 데이터/명령"],
            ["태블릿 앱 → 백엔드", "REST :8002 (HTML+fetch)", "수동배차/상태/제어"],
            ["백엔드 → 로봇", "REST :8090", "이동/잭/맵/설정 명령"],
            ["백엔드 → 로봇", "WebSocket :8090/ws/v2/topics", "토픽 구독(planning_state, jack_state, detected_rack, nearby_robots)"],
            ["백엔드 → 관제웹", "폴링 응답", "SSE/WS 미사용 — 1~2초 주기 polling"],
        ],
        col_widths=[4.5, 4.5, 7],
    )

    para(
        doc,
        "※ 의도적으로 푸시(SSE/WS) 채널을 두지 않음. 단순화·디버깅 용이성·다중 로봇 환경에서의 "
        "타이밍 명확성을 위해 폴링 기반으로 통일.",
        color=MUTED,
        size=9.5,
    )

    # ───────────────────────────────────────────────
    h1(doc, "5. 데이터 모델 (MySQL rcs_basic_db)")
    para(doc, "주요 테이블과 핵심 컬럼만 발췌. 외래키/타임스탬프 등은 생략.", color=MUTED, size=9.5)
    table(
        doc,
        ["테이블", "핵심 컬럼"],
        [
            ["businesses", "사업장 정보"],
            ["areas", "영역 — is_main_floor (메인층 / 잭 든 채로 작업할 층)"],
            ["robots", "ip_address, model, robot_type(lifting/serving), charging_id, standby_id, max_speed"],
            ["robot_status", "실시간 상태 1행: battery_level, position_x/y/yaw, status(0~4)"],
            ["robot_status_history", "시계열 이력"],
            ["robot_maps", "매핑 산출물: bag_url, image_url, grid_origin_*, robot_map_id(펌웨어 맵 ID)"],
            ["map_pois", "x/y, world_x/y, angle, poi_type, rack_size(S300/S600), area_name"],
            ["map_lines", "from_poi_id, to_poi_id, direction(forward/backward/bidirectional), line_type"],
            ["map_polygons", "shape_type='polygon' | 'firewall' | 'zone', points_json"],
            ["task_routes", "name, work_mode(rack_pickup/delivery_no_rack/simple_move)"],
            ["task_route_waypoints", "route_id, poi_id, order, waypoint_type, wait_sec"],
            ["scheduled_tasks", "robot_id, route_id, start/end_time, repeat_type, repeat_days, start/end_date"],
            ["task_history", "task_id, robot_id, status, started_at, finished_at, error_message"],
            ["activity_logs", "category, action, message, source — 운영 감사 로그"],
            ["alarm_logs", "배터리/오류/이탈 등 알람"],
            ["users", "단순 인증"],
        ],
        col_widths=[4, 12],
    )

    # ───────────────────────────────────────────────
    h1(doc, "6. 핵심 기능 — 라이다 기반 랙 인식 및 잭킹")
    para(
        doc,
        "본 시스템의 가장 정교한 동작은 \"운반할 랙(카트) 아래로 정확히 들어가서 잭으로 들어 올리고, "
        "목적지에 정확한 자세로 내려놓는\" 잭킹 흐름이다. 이를 위해 AutoXing 펌웨어의 LiDAR 기반 "
        "랙 인식 알고리즘과 관제측의 사전 등록 정보(rack.specs, Shelves Point)가 함께 동작한다.",
    )

    h2(doc, "6.1 Shelves Point (필수 등록 요건)")
    para(
        doc,
        "align_with_rack 명령은 단순한 좌표 이동이 아니라 \"등록된 랙 위치에서 LiDAR 로 랙 다리를 "
        "감지하면서 중앙 정렬한다\"는 정밀 모드이다. 이 동작을 트리거하려면 펌웨어 맵의 overlay 에 "
        "POI 가 type=\"34\", subtype=\"rack\" 으로 등록되어 있어야 한다. 일반 type 으로는 "
        "\"failed to find rack\" 오류로 종료된다.",
    )
    code_block(
        doc,
        '{\n'
        '  "type": "34",\n'
        '  "subtype": "rack",\n'
        '  "shelvesState": "0",\n'
        '  "hasFixedLegs": false,\n'
        '  "dockViaDirection": "front",\n'
        '  "mapOverlay": true\n'
        '}',
    )
    para(
        doc,
        "맵 저장/동기화 시 시스템은 jack 타입 POI 들을 위 형식으로 자동 변환해 overlay 에 포함시킨다. "
        "중복 등록을 막기 위해 기존 type=\"34\" 항목을 모두 제거한 뒤 새로 추가한다.",
    )

    h2(doc, "6.2 rack.specs — 랙 사이즈 매칭")
    para(
        doc,
        "펌웨어는 등록된 rack.specs 후보들 중 LiDAR 측정 다리 패턴과 가장 가까운 spec 을 매칭한다. "
        "여러 spec 을 등록하면 잘못된 spec 과 매칭되어 \"Wrong rack size\"(12006) 오류가 빈번해진다. "
        "이를 회피하기 위해 본 시스템은 로봇 모델명 기준 단일 spec 만 동기화한다.",
    )
    h3(doc, "S300 (소형) 정의")
    code_block(
        doc,
        '{"width": 0.765, "depth": 0.765,\n'
        ' "margin": [0.0925]*4,\n'
        ' "alignment": "center", "alignment_margin_back": 0.02,\n'
        ' "leg_shape": "other", "leg_size": 0.04,\n'
        ' "foot_radius": 0.02,\n'
        ' "cargo_to_jack_front_edge_min_distance": 0.05}',
    )
    h3(doc, "S600 (대형) 정의")
    code_block(
        doc,
        '{"width": 0.83, "depth": 0.87,\n'
        ' "margin": [0.1]*4,\n'
        ' "alignment": "center", "alignment_margin_back": 0.02,\n'
        ' "leg_shape": "other", "leg_size": 0.05,\n'
        ' "foot_radius": 0.025,\n'
        ' "cargo_to_jack_front_edge_min_distance": 0.05}',
    )
    h3(doc, "로봇 모델 → 단일 spec 매핑")
    table(
        doc,
        ["robot.model", "사용 spec"],
        [["crawler_s300_op5", "S300"], ["crawler_heavy", "S600"], ["(매핑 없음)", "S300 (안전 디폴트)"]],
        col_widths=[7, 9],
    )

    h2(doc, "6.3 align_with_rack 정밀 정렬 시퀀스")
    para(doc, "코드 위치: BackEnd/app/services/jack_service.py — align_with_retry()", color=MUTED, size=9.5)
    para(doc, "다음 시퀀스로 진행한다.", bold=True)
    bullet(doc, "[1] 등록된 Shelves Point 좌표로 align_with_rack 명령 — 펌웨어가 path planner 로 접근")
    bullet(doc, "[2] LiDAR 로 랙 다리 4점 패턴 인식 (rack.specs 매칭)")
    bullet(doc, "[3] 다리 사이 중심점 계산 → 미세 조정 이동으로 본체 정렬")
    bullet(doc, "[4] 도착 시 succeeded 반환 — 그 직후 호출측이 jack_up 호출하면 랙 들기 완료")
    para(doc, "실패 케이스 (rack_detection_error) 4단계 자동 회복:", bold=True)
    table(
        doc,
        ["시도", "조치", "의미"],
        [
            ["1회차", "그대로 재실행", "일시적 LiDAR 노이즈 회피"],
            ["2회차", "위치 재로컬화 (start_global_positioning) 후 재시도", "로컬화 드리프트 보정"],
            ["3회차", "0.5m 후진 + 재로컬화 후 재시도", "랙 너무 가까이서 시작한 경우 회피"],
            ["4회차", "좌측 30cm 측면 우회 + 재로컬화 후 재시도", "특정 각도에서 다리 잘림 회피"],
        ],
        col_widths=[2.5, 6, 7.5],
    )
    para(
        doc,
        "4회 모두 실패 시 작업은 \"랙 정렬 실패\" 로 종료되고 activity_logs / task_history 에 "
        "기록된다. 운영자가 수동으로 \"현 위치\" 버튼으로 POI 좌표를 재등록하면 대부분 해결된다.",
    )

    h2(doc, "6.4 잭킹 흐름 (rack_pickup 모드)")
    code_block(
        doc,
        "충전소(C1)\n"
        "  → 사전 접근 POI (운영자가 'C1-1' 등록 시 자동 경유) 또는 충전소 좌표 직접\n"
        "  → standby POI (R1, 랙 위치) 까지 standard 이동\n"
        "  → align_with_rack (LiDAR 랙 다리 인식 → 중앙 정렬)\n"
        "  → jack_up (10초 대기, 푸트프린트가 랙 사이즈로 확장됨 — /robot_model 토픽으로 반영)\n"
        "  → to_unload_point (정밀 이동, '랙 든 상태' path)\n"
        "  → jack_down (10초 대기 + 400 오류 시 5초 간격 5회 재시도)\n"
        "  → 다음 jack POI 가 있으면 다시 align → jack_up → to_unload_point → jack_down ... 반복\n"
        "  → 마지막 jack POI 처리 후 R1 으로 복귀 (잭 상태에 따라 align+jack_up 후 이동)\n"
        "  → R1 에서 jack_down 으로 랙 보관\n"
        "  → 충전소(C1) 복귀 — 사전 접근 POI (C1-1) → charge 명령 (5회 재시도)",
    )

    h2(doc, "6.5 충전소 도킹 (자동 + 수동 통일)")
    para(
        doc,
        "충전소 POI(예: C1) 옆에 같은 맵에 \"<C1>-1\" 이름 POI 를 등록해 두면, 시스템이 자동으로 "
        "해당 위치로 먼저 standard 이동 후 charge 명령을 실행한다. 등록 안 된 경우 충전소 좌표로 "
        "직접 standard 이동 후 charge — 동작은 동일.",
    )
    bullet(doc, "1단계: standard 이동 (사전 접근 POI 또는 충전소 좌표)")
    bullet(doc, "2단계: 안정화 sleep 2초")
    bullet(doc, "3단계: charge 명령 (cx, cy, cyaw, charge_retry_count=3) — 5회 외부 재시도 + 각 120초 timeout")
    para(
        doc,
        "적용 위치 4곳 모두 동일:\n"
        " (1) scheduler._return_to_charger (작업 종료 후 자동)\n"
        " (2) jack_service charging waypoint (경로 중 충전 단계)\n"
        " (3) robot.api_dock_to_charger (제어패널 dock 버튼)\n"
        " (4) jack_service.force_return_and_dock (강제 종료 시 — (1) 호출)",
    )

    # ───────────────────────────────────────────────
    h1(doc, "7. 작업 모드 (work_mode) 3종 비교")
    table(
        doc,
        ["항목", "rack_pickup", "delivery_no_rack", "simple_move"],
        [
            ["잭업 의미", "랙(카트)을 들어 올림", "잭 상판 적재물 들어 올림", "사용 안 함"],
            ["align_with_rack 필요", "예 (랙 다리 인식)", "아니오", "아니오"],
            ["이동 명령 type", "to_unload_point (정밀)", "standard", "standard"],
            ["허용 로봇 타입", "lifting", "lifting", "lifting, serving"],
            ["대표 시나리오", "공장 카트 셔틀", "물품 배달", "안내/서빙 단순 이동"],
            ["추가 목적지 사이클", "align+잭업+이동+잭다운", "잭업+이동+잭다운", "이동만"],
        ],
        col_widths=[4, 4, 4, 4],
    )

    # ───────────────────────────────────────────────
    h1(doc, "8. 다중 로봇 조율 (충돌·데드락 회피)")

    h2(doc, "8.1 POI 락 (poi_lock.py)")
    para(
        doc,
        "정적 목적지(잭/대기/충전소 POI)를 같은 시점에 두 로봇이 점유하지 못하도록 in-memory 락. "
        "작업 시작 시 경로상 모든 POI 를 try_acquire 로 한 번에 락 — 하나라도 충돌하면 전체 롤백. "
        "작업 종료/예외/취소 시 release_all_by_robot 로 일괄 해제. 같은 로봇은 재진입 허용.",
    )

    h2(doc, "8.2 Zone 락 (zone_lock.py + zone_guard.py)")
    para(
        doc,
        "좁은 양방향 통로 등 \"한 번에 한 로봇만 통과해야 하는 구간\" 을 MapPolygon 의 "
        "shape_type='zone' 으로 등록한다. 이동 직전 zone_guard 가 현재 pose → 목적지 선분이 "
        "어떤 zone 폴리곤과 교차/포함되는지 geometry.py 의 ray casting + 선분 교차 알고리즘으로 판정. "
        "필요한 zone 모두 락 획득까지 최대 60초 대기.",
    )
    bullet(doc, "load_zone_polygons(map_id) — 해당 맵의 모든 zone 폴리곤 로드")
    bullet(doc, "zones_crossed_by_segment(zones, p1, p2) — 선분이 통과할 zone id 리스트")
    bullet(doc, "zones_containing_point(zones, x, y) — 끝점이 포함된 zone id 리스트")
    bullet(doc, "try_acquire(zone_ids, robot_id) — 원자적 락 시도, 충돌 시 전체 롤백")

    h2(doc, "8.3 Deadlock Monitor (deadlock_monitor.py) — 현재 비활성")
    para(
        doc,
        "AutoXing 펌웨어는 P2P 분산 회피만 수행해 좁은 통로에서 두 로봇이 마주보거나 동일 목적지를 "
        "점유하면 풀리지 않는 정체가 생긴다. 본 모니터는 WebSocket 으로 /planning_state, /nearby_robots "
        "를 5초 주기 폴링하며 다음 두 패턴을 25초 연속 관측 시 데드락으로 판정하고 한쪽을 양보시킨다.",
    )
    bullet(doc, "(A) is_waiting_for_dest == True 가 연속 5샘플 (25초)")
    bullet(doc, "(B) move_state == 'moving' 인데 remaining_distance 변화가 0.05m 미만")
    bullet(doc, "근접 로봇 감지(nearby_count > 0)가 3샘플 이상 지속되면 임계 단축 (15초)")
    bullet(doc, "양보 대상 = 데드락 후보 중 robot_id 가장 큰 로봇")
    bullet(doc, "양보 후 30초 대기 → POI 락 재획득 가능하면 별도 스레드로 재실행")
    bullet(doc, "정밀 동작(align/잭킹/충전/대기 상태) 은 정상 미세 움직임이므로 감지 제외")
    para(
        doc,
        "현재 운영 사이트엔 좁은 통로가 없어 main.py 에서 비활성(주석) 처리. POI 락 / Zone 락 / "
        "safe_move 자동 재시도 / align_with_retry 4단계 회복은 그대로 동작.",
        color=MUTED, size=9.5,
    )

    h2(doc, "8.4 자동 재시도용 메타데이터")
    bullet(doc, "_active_route_jobs[ip] — 진행 중 작업의 입력(waypoints, manual_confirm, area_id, work_mode) 보관")
    bullet(doc, "_paused_route_jobs[ip] — 양보로 일시정지된 작업 + paused_at + retries_left")
    bullet(doc, "consume_paused_job(ip) — 재실행 직전 메타 꺼내가며 제거")

    h2(doc, "8.5 safe_thread (thread_utils.py)")
    para(
        doc,
        "daemon 백그라운드 스레드의 예외가 조용히 사라지지 않도록 try/except 로 래핑하는 헬퍼. "
        "예외 발생 시 logger.exception + activity_logs(category='system', action='thread_crash') 에 기록.",
    )

    # ───────────────────────────────────────────────
    h1(doc, "9. 사용자 제어 — 정지/일시정지/강제종료")
    table(
        doc,
        ["명령", "API", "동작 요약"],
        [
            ["일시정지(pause)", "POST /api/robots/remote/pause/{ip}",
             "_paused_flags set + 현재 이동 cancel. 재개 시 같은 이동 자동 재시도."],
            ["재개(resume)", "POST /api/robots/remote/resume/{ip}", "paused 해제, 대기 중인 wait_move/safe_move/_interruptible_sleep 진행"],
            ["즉시 중단", "POST /api/robots/remote/stop-all/{ip}",
             "_stop_flags set + 이동 cancel + jack_down. 그 자리 정지(자동 복귀 X)."],
            ["강제 종료", "POST /api/robots/remote/force-return/{ip}",
             "work_mode 별 분기 (아래 표)"],
            ["수동 confirm", "POST /api/robots/remote/confirm/{ip}", "wait_for_confirm 대기 즉시 깨움 → 출발"],
            ["다음 포인트", "POST /api/robots/remote/next-point/{ip}", "waiting_next_or_return 상태에서 추가 목적지 지정"],
            ["복귀", "POST /api/robots/remote/return/{ip}", "waiting_next_or_return 상태에서 복귀 트리거"],
        ],
        col_widths=[3, 6, 7],
    )

    h2(doc, "9.1 강제 종료 work_mode 분기")
    table(
        doc,
        ["work_mode", "동작"],
        [
            ["rack_pickup", "잭업(이미 들고 있어도 시도) → R1(랙 위치) 복귀 → 잭다운 → 충전소"],
            ["delivery_no_rack", "안전 위해 잭다운 한 번 시도 → 충전소"],
            ["simple_move", "잭 조작 없이 곧장 충전소"],
        ],
        col_widths=[4, 12],
    )
    para(
        doc,
        "프론트엔드 confirm 문구도 work_mode 에 따라 자동 분기 (\"랙을 원래 위치(R1)에 두고 ...\", "
        "\"잭을 내린 뒤 충전소로 복귀\", \"곧바로 충전소로 복귀\").",
    )

    h2(doc, "9.2 좀비 스레드 방지 — stop_robot_job 강화")
    para(
        doc,
        "초기 구현에서는 stop_robot_job 이 _stop_flags 만 set 하므로 wait_for_confirm 안에서 "
        "Event.wait(30분) 으로 대기 중인 작업은 즉시 빠져나오지 못해 최대 30분 좀비 thread 가 발생했다. "
        "현재는 stop_robot_job 이 _confirm_events 의 Event 도 set 하여 wait_for_confirm 이 즉시 깨어나고, "
        "곧이은 _check_stop() 에서 RuntimeError 로 정상 종료된다.",
    )
    para(
        doc,
        "wait_for_confirm 자체도 try/finally 로 _confirm_events 정리 보장 — 예외 시 좀비 entry 누적 방지.",
    )

    # ───────────────────────────────────────────────
    h1(doc, "10. 백엔드 API 카테고리")

    h2(doc, "10.1 인증 / 사용자")
    bullet(doc, "POST /api/auth/login, POST /api/auth/logout")
    bullet(doc, "GET/PUT /api/users/me/password")

    h2(doc, "10.2 로봇 관리 / 실시간")
    bullet(doc, "GET /api/robots — 목록")
    bullet(doc, "POST /api/robots/register — IP 입력 → 자동 정보 조회 후 등록")
    bullet(doc, "GET/PATCH/DELETE /api/robots/{id}")
    bullet(doc, "POST /api/robots/{id}/switch-floor — 층 전환 + 맵 동기화")
    bullet(doc, "GET /api/robots/quick-status/{ip} — 배터리, run_state, 온라인")
    bullet(doc, "GET /api/robots/job-status — 모든 로봇 작업 상태 (work_mode 노출)")
    bullet(doc, "GET /api/robots/job-status/{ip}")

    h2(doc, "10.3 로봇 원격 제어")
    bullet(doc, "POST /api/robots/remote/{pause|resume|stop-all|confirm|return|next-point|force-return}/{ip}")
    bullet(doc, "POST /api/robots/remote/dock/{ip} — 충전소 복귀 (C1-1 자동 경유)")
    bullet(doc, "POST /api/robots/remote/jack/{ip}/{jack_up|jack_down}")
    bullet(doc, "POST /api/robots/remote/cancel-move/{ip}")
    bullet(doc, "POST /api/robots/remote/return-to-standby/{ip} — 잭업 → R1 → 잭다운")
    bullet(doc, "POST /api/robots/remote/relocalize/{ip} — 위치 복구")
    bullet(doc, "POST /api/robots/remote/control-mode/{ip} — auto/remote 모드 전환")
    bullet(doc, "POST /api/robots/remote/twist/{ip} — Twist (linear/angular velocity)")
    bullet(doc, "POST /api/robots/remote/shutdown/{ip} — 로봇 종료")

    h2(doc, "10.4 맵 / POI")
    bullet(doc, "GET /api/map/active-pois, GET /api/map/default-area")
    bullet(doc, "맵 매핑 시작/중단/저장, sync (일반 / full restart)")
    bullet(doc, "POI / 라인 / 폴리곤 CRUD — 이름 중복 사전 체크")
    bullet(doc, "Shelves Point 자동 포함 (rack_pickup 동작 필수)")
    bullet(doc, "rack.specs 자동 단일화 동기화")

    h2(doc, "10.5 작업 관리")
    bullet(doc, "경로(routes) CRUD + work_mode 검증 + 이름 중복 체크")
    bullet(doc, "스케줄 CRUD (once/daily/weekly), 즉시 실행, 토글, 종료 시간까지 반복")
    bullet(doc, "POST /api/tasks/manual-run-pois — 수동 배차 (work_mode 전달)")
    bullet(doc, "POST /api/tasks/batch-manual — 다중 로봇 배치 배차 + 반복 횟수")
    bullet(doc, "이력 조회 (필터/통계)")
    bullet(doc, "태블릿 페이지: GET /api/tasks/tablet/{robot_id}, GET /api/tasks/tablet-simple/{robot_id}")

    h2(doc, "10.6 설정 / 로그 / 백업")
    bullet(doc, "GET/PATCH /api/settings/system-name — RCS·태블릿 공통 시스템 이름 SSOT")
    bullet(doc, "GET /api/activity-logs — 카테고리/액션/소스 필터, 페이징")
    bullet(doc, "POST /api/backup/create, POST /api/backup/restore — DB 백업")
    bullet(doc, "잭 테스트 — POST /api/jack-test/start, GET /api/jack-test/status/{job_id}")
    bullet(doc, "헬스 체크 — GET /ping, GET /health (DB 연결 확인)")

    # ───────────────────────────────────────────────
    h1(doc, "11. 관제 웹 프론트엔드 기능")

    h2(doc, "11.1 모니터링 페이지 (3D)")
    bullet(doc, "Three.js 로 맵 이미지를 텍스처화 + 로봇/POI/라인/폴리곤 3D 마커")
    bullet(doc, "로봇 풋프린트 박스 — 잭 업 시 랙 사이즈(0.83 x 0.87 등)로 확장 시각화")
    bullet(doc, "잭 POI: 2D 박스(rack 크기) + V자 쉐브론(접근 방향) + 90도 보정")
    bullet(doc, "충전소 마커는 도킹 자세 반대 0.5m 이동 — 로봇 도킹 시 가려지지 않게")
    bullet(doc, "좌측 패널: 수동 배차 + 작업 중 로봇 카드 (ActiveJobsPanel)")
    bullet(doc, "  ⏸ 정지 / ▶ 재개 / ■ 강제 종료 (work_mode 별 confirm 문구)")
    bullet(doc, "우측 패널: 작업 현황 (JobStatusPanel) + 진행 카드 + 오늘 스케줄 + 최근 이력")
    bullet(doc, "로봇 페이지네이션 (5대 단위), 색상별 이동 경로 표시 (현재→목적지 직선)")
    bullet(doc, "잭 테스트 패널, 데이터 로딩 스피너")

    h2(doc, "11.2 맵 페이지")
    bullet(doc, "매핑 시작/중단/저장 — 영역 이름 중복 사전 체크")
    bullet(doc, "POI 편집 팝업 — 타입별 prefix 자동 생성 (C 충전소, J 작업, R 랙, W 경유지)")
    bullet(doc, "작업 포인트(jack POI) 생성 — 보라색 마커")
    bullet(doc, "가상벽(폴리곤) / zone 폴리곤 그리기")
    bullet(doc, "충전소 자동 등록, \"현 위치\" 버튼으로 POI 재등록")
    bullet(doc, "맵 ↔ 모니터링 동기화 (default-area 기반)")
    bullet(doc, "맵 저장 후 자동 sync — 소스 로봇 위치 보정 + grid_origin 자동 보정")

    h2(doc, "11.3 로봇 관리")
    bullet(doc, "IP 입력 → 펌웨어 직접 조회 → 정보 자동 등록")
    bullet(doc, "등록 즉시 리스트 갱신")
    bullet(doc, "robot_type (lifting/serving) 지정")
    bullet(doc, "max_speed DB 저장 + 서버 시작 시 자동 적용 (apply_saved_speeds)")
    bullet(doc, "로봇별 charging_id / standby_id 지정")
    bullet(doc, "설정 완료 모달")

    h2(doc, "11.4 작업 관리")
    bullet(doc, "경로 카드 UI — 드래그 순서 변경, 타입 선택, 대기시간")
    bullet(doc, "스케줄 — CLOi 스타일 모달 (날짜/시간/반복 요일/종료)")
    bullet(doc, "수동 배차 + 일괄 배차 (다중 로봇 동시 시작)")
    bullet(doc, "이름 중복 체크, 커스텀 alert/confirm 모달")

    h2(doc, "11.5 설정 / 로그")
    bullet(doc, "비밀번호 변경, DB 백업/복원")
    bullet(doc, "시스템 이름 변경 — 서버 단일 진실 원천. TopBar 마운트 시 fetch + custom event 즉시 갱신")
    bullet(doc, "활동 로그 조회 (필터/페이징)")

    # ───────────────────────────────────────────────
    h1(doc, "12. 태블릿 앱")

    h2(doc, "12.1 TabletApp (기존, 수동배차 + 제어)")
    bullet(doc, "헤더: 시스템 이름(60초 동기화) + 로봇 이름 + 상태 + 종료/설정/홈 버튼")
    bullet(doc, "수동 배차: 작업 종류(work_mode) 선택 + 픽업/드롭오프 (lifting → 3종, serving → simple_move 만)")
    bullet(doc, "simple_move 선택 시 라벨 자동 변경 (\"시작 위치 / 도착 위치\")")
    bullet(doc, "진행 상태: 작업 종류, 현재 단계, 출발 대기 시 5초 자동 confirm")
    bullet(doc, "제어 패널 모달 — 잭 업/다운, 이동 취소, 충전소 복귀, 원격 조종, 위치 복구")
    bullet(doc, "스케줄 목록 + 즉시 실행/토글, 시간 범위 밖 비활성")
    bullet(doc, "장애물 감지 시 배너 + 알림음")

    h2(doc, "12.2 TabletAppSimple (단순 표시)")
    bullet(doc, "좌상단: 앱 네이티브 설정 버튼 (어두운 원형, elevation — WebView 위 가시성 보장)")
    bullet(doc, "상단 가운데: 시스템 이름 (60초 동기화)")
    bullet(doc, "우상단: 로봇 이름")
    bullet(doc, "중앙 가장 큰 글씨: 현재 상태 (이동중/대기/충전 등 — 상태별 색상)")
    bullet(doc, "그 밑: 목적지 이름")
    bullet(doc, "오프라인 배너 (job-status fetch 실패 시)")

    h2(doc, "12.3 WebView 동작")
    para(
        doc,
        "두 앱 모두 Android 측은 단순 WebView shell 이고, 화면은 서버가 동적 렌더링한 HTML 을 "
        "WebView 가 로드한다. (수정 시 APK 재빌드 불필요 — 서버 재배포만)",
    )
    bullet(doc, "TabletApp: WebView.loadUrl(\"http://서버IP:8002/api/tasks/tablet/{robot_id}\")")
    bullet(doc, "TabletAppSimple: WebView.loadUrl(\"http://서버IP:8002/api/tasks/tablet-simple/{robot_id}\")")
    bullet(doc, "첫 실행 시 다이얼로그 — 서버 주소 / 로봇 ID 입력 → SharedPreferences 저장")
    bullet(doc, "JavascriptInterface (\"Android\") — openSettings / goHome 브리지")
    bullet(doc, "FLAG_KEEP_SCREEN_ON, hideSystemUI — 키오스크 모드")

    # ───────────────────────────────────────────────
    h1(doc, "13. AutoXing 로봇 API 핵심")
    table(
        doc,
        ["용도", "메서드 / 경로"],
        [
            ["이동 명령", "POST /chassis/moves (type: standard, align_with_rack, to_unload_point, charge)"],
            ["이동 조회/취소", "GET /chassis/moves/{id} / PATCH /chassis/moves/current {state:'cancelled'}"],
            ["현재 pose", "GET /chassis/pose"],
            ["잭", "POST /services/jack_up, POST /services/jack_down"],
            ["맵", "GET/POST/PATCH/DELETE /maps/{id}"],
            ["현재 맵 선택", "POST /chassis/current-map (overlays_version 강제 reload)"],
            ["사용자 설정", "GET/PATCH /system/settings/user (rack.specs 포함)"],
            ["서비스 재시작", "POST /services/restart_py_axbot (full sync 후)"],
            ["위치 보정", "POST /services/start_global_positioning"],
            ["WebSocket", "ws://{ip}:8090/ws/v2/topics (enable_topic 메시지로 구독)"],
        ],
        col_widths=[3.5, 12.5],
    )

    h3(doc, "핵심 노하우")
    bullet(doc, "overlays_version 자동 갱신 안 됨 → PATCH 후 POST /chassis/current-map 으로 강제 reload")
    bullet(doc, "full 동기화 흐름: 서비스 재시작 → 90초 대기 → Shelves Point PATCH → current-map 재선택")
    bullet(doc, "잭 다운 직후 짧은 시간 동안 명령에 400 반환 — 5초 간격 5회 자동 재시도")
    bullet(doc, "create_move 실패 (failed to calc global path 등) → safe_move 가 자동 재시도")

    # ───────────────────────────────────────────────
    h1(doc, "14. 운영 / 배포")
    h2(doc, "14.1 도커 구성")
    code_block(
        doc,
        "services:\n"
        "  backend:\n"
        "    build: ./BackEnd\n"
        "    ports: [\"8002:8002\"]\n"
        "    environment:\n"
        "      - DB_HOST=${DB_HOST:-host.docker.internal}\n"
        "      - DB_USER=${DB_USER}\n"
        "      - DB_PASSWORD=${DB_PASSWORD}\n"
        "      - DB_NAME=${DB_NAME}\n"
        "    extra_hosts:\n"
        "      - \"host.docker.internal:host-gateway\"\n"
        "    healthcheck:\n"
        "      test: [\"CMD\", \"python\", \"-c\", \"import urllib.request; urllib.request.urlopen('http://localhost:8002/ping')\"]\n"
        "\n"
        "  frontend:\n"
        "    build:\n"
        "      context: ./frontend\n"
        "      args:\n"
        "        - NEXT_PUBLIC_API_URL=${NEXT_PUBLIC_API_URL}\n"
        "    ports: [\"3002:3002\"]\n"
        "    depends_on:\n"
        "      backend: { condition: service_healthy }",
    )

    h2(doc, "14.2 .env 예시")
    code_block(
        doc,
        "SERVER_IP=192.168.20.3\n"
        "DB_HOST=host.docker.internal\n"
        "DB_PORT=3306\n"
        "DB_USER=root\n"
        "DB_PASSWORD=1234\n"
        "DB_NAME=rcs_basic_db\n"
        "NEXT_PUBLIC_API_URL=http://192.168.20.3:8002",
    )

    h2(doc, "14.3 배포 절차")
    bullet(doc, "git pull origin feature/backend_noah")
    bullet(doc, "DB 컬럼 마이그레이션 (ALTER TABLE; 새 컬럼 robot_type / work_mode / rack_size / area_name)")
    bullet(doc, "docker compose down --remove-orphans")
    bullet(doc, "docker compose build --no-cache frontend backend")
    bullet(doc, "docker compose up -d")
    bullet(doc, "curl /ping, /health, /api/settings/system-name 으로 검증")

    h2(doc, "14.4 IP 변경 시")
    bullet(doc, ".env 의 SERVER_IP / NEXT_PUBLIC_API_URL 갱신")
    bullet(doc, "프론트 캐시 무시 재빌드 (NEXT_PUBLIC_API_URL 빌드 시 인라인)")
    bullet(doc, "태블릿 앱 설정에서 서버 주소 변경")

    h2(doc, "14.5 접속 주소")
    table(
        doc,
        ["용도", "URL"],
        [
            ["관제 웹", "http://{서버IP}:3002"],
            ["API", "http://{서버IP}:8002"],
            ["태블릿 (수동배차+제어)", "http://{서버IP}:8002/api/tasks/tablet/{robot_id}"],
            ["태블릿 (단순)", "http://{서버IP}:8002/api/tasks/tablet-simple/{robot_id}"],
            ["헬스체크", "http://{서버IP}:8002/health"],
        ],
        col_widths=[5, 11],
    )

    # ───────────────────────────────────────────────
    h1(doc, "15. 알려진 미해결 이슈")
    bullet(doc, "align_with_rack 실패(rack_detection_error) — 4단계 회복으로도 인식 못 하는 케이스 → 운영자 \"현 위치\" 재등록 필요")
    bullet(doc, "rack_area_id 사용 불가 (regionType 미확인 — AutoXing 문의 필요)")
    bullet(doc, "detectRackSize REST 없음 (SDK 전용)")
    bullet(doc, "잭 다운 후 로봇 빠져나오기 시간 불확실 (고정 10초 + 400 시 5초 간격 재시도로 보정)")
    bullet(doc, "to_unload_point J1 이동 미작동 사례 — 좌표 정확도 또는 맵 grid_origin 의심")
    bullet(doc, "맵 변경 시 경로 waypoint POI ID 자동 재매핑 미구현 (현재 수동 DB UPDATE)")

    # ───────────────────────────────────────────────
    h1(doc, "16. 보안 / 품질 특징")
    bullet(doc, "인증: 단순 password (관리자급 단일 계정)")
    bullet(doc, "HTTPS 미적용 — 로컬망 사용 전제")
    bullet(doc, "WebSocket 보안 없음 — 로봇이 신뢰망 안에 있다는 전제")
    bullet(doc, "스레드 관리: safe_thread 도입 — 일부 라우터는 직접 threading.Thread 사용 (점진 마이그레이션 중)")
    bullet(doc, "wait_for_confirm: stop_robot_job 이 confirm event 까지 set → 즉시 깨움 (좀비 thread 방지)")
    bullet(doc, "DB session: try/finally 로 close 보장, get_db dependency 가 rollback 처리")
    bullet(doc, "백업: DB dump 라우터 + activity_logs 함께 백업")

    # ───────────────────────────────────────────────
    h1(doc, "부록 A. 주요 코드 위치")
    table(
        doc,
        ["영역", "파일"],
        [
            ["FastAPI 진입점", "BackEnd/app/main.py"],
            ["DB 연결", "BackEnd/app/database.py"],
            ["로봇 라우터", "BackEnd/app/routers/robot.py"],
            ["맵 라우터", "BackEnd/app/routers/map.py"],
            ["작업 라우터 (태블릿 페이지 포함)", "BackEnd/app/routers/task.py"],
            ["설정 라우터 (system-name)", "BackEnd/app/routers/settings.py"],
            ["경로 실행 / 강제종료", "BackEnd/app/services/jack_service.py"],
            ["스케줄러 / 충전소 자동 복귀", "BackEnd/app/services/scheduler.py"],
            ["POI 락", "BackEnd/app/services/poi_lock.py"],
            ["Zone 락 / Zone Guard", "BackEnd/app/services/zone_lock.py, zone_guard.py"],
            ["폴리곤 지오메트리", "BackEnd/app/services/geometry.py"],
            ["데드락 모니터", "BackEnd/app/services/deadlock_monitor.py"],
            ["safe_thread", "BackEnd/app/services/thread_utils.py"],
            ["rack.specs / robot_types 상수", "BackEnd/app/constants/"],
            ["태블릿 템플릿", "BackEnd/app/templates/tablet.html, tablet_simple.html"],
            ["프론트 모니터링", "frontend/app/monitoring/"],
            ["프론트 컴포넌트(작업 카드)", "frontend/app/components/ui/monitoring/"],
            ["Android (수동배차+제어)", "TabletApp/"],
            ["Android (단순 표시)", "TabletAppSimple/"],
        ],
        col_widths=[6, 10],
    )

    # 푸터
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— 끝 —")
    _set_font(r, size=10, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"saved: {OUTPUT}")
    print(f"size : {OUTPUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    build()
